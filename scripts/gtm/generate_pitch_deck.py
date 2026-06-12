#!/usr/bin/env python3
"""Generate Tier 0 pitch deck PPTX (PL + EN) from one-pager content."""

from __future__ import annotations

from pathlib import Path

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
except ImportError as exc:
    raise SystemExit("Install python-pptx: pip install python-pptx") from exc

REPO = Path(__file__).resolve().parents[2]
OUT_PL = REPO / "docs" / "gtm" / "pl" / "PITCH_DECK.pptx"
OUT_EN = REPO / "docs" / "gtm" / "en" / "PITCH_DECK.pptx"


def add_title_slide(prs: Presentation, title: str, subtitle: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    slide.placeholders[1].text = subtitle


def add_bullet_slide(prs: Presentation, title: str, bullets: list[str]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    body = slide.placeholders[1].text_frame
    body.clear()
    for i, line in enumerate(bullets):
        p = body.paragraphs[0] if i == 0 else body.add_paragraph()
        p.text = line
        p.level = 0
        p.font.size = Pt(20)


def build_deck(lang: str) -> Presentation:
    prs = Presentation()
    if lang == "pl":
        add_title_slide(prs, "4VELO / SPORT", "Platforma sportowo-rywalizacyjna B2B/B2G — v1.0")
        add_bullet_slide(prs, "Problem rynku", [
            "Oszustwa w rankingach bez egzekucji zasad",
            "Niestabilny GPS w tle i awarie przy szczytach (~150k sesji u konkurencji)",
            "Decydenci (JST, HR, kluby) nie ufają wynikom",
        ])
        add_bullet_slide(prs, "Specjalizacja: 3 dyscypliny", [
            "Rower · Bieg · Nordic walking",
            "GPS, anti-cheat i scoring pod te aktywności",
            "Bez katalogu 18 dyscyplin i wellness tylko na krokach",
        ])
        add_bullet_slide(prs, "Pozycja 4VELO", [
            "White-label multi-tenant",
            "Anti-cheat 4 warstwy + publiczna polityka trust",
            "Odporność mobilna (outbox) + backend scale-tested 10k–300k",
            "Konsola operatora + ligi 1v1 w API",
        ])
        add_bullet_slide(prs, "Zaufanie i skala", [
            "Publiczna Anti-Cheat / Scoring Policy (URL live)",
            "Runbook RSP-weekend + raport symulacji 10k–300k",
            "DPIA + lista podprocesorów per tenant",
        ])
        add_bullet_slide(prs, "3 ścieżki przychodu", [
            "Miasta / JST — kampanie sezonowe",
            "Firmy — ligi działów / firma vs firma (HR)",
            "Kluby — klub vs klub, sezony federacyjne",
        ])
        add_bullet_slide(prs, "Piloty dowodu (90 dni)", [
            "Pilot A: 4 działy, 6 tygodni, 1 tenant",
            "Pilot B: 2 kluby, 30 dni, 1v1",
            "Następnie: liga wielofirmowa + case study",
        ])
        add_title_slide(prs, "Następny krok", "Demo 30 min · DPIA · start pilota w 2 tygodnie")
    else:
        add_title_slide(prs, "4VELO / SPORT", "B2B/B2G sports competition platform — v1.0")
        add_bullet_slide(prs, "Market problem", [
            "Ranking cheating with no rule enforcement",
            "Unstable background GPS and peak crashes (~150k sessions at competitors)",
            "Buyers (local gov, HR, clubs) do not trust results",
        ])
        add_bullet_slide(prs, "Specialization: 3 disciplines", [
            "Cycling · Running · Nordic walking",
            "GPS, anti-cheat and scoring built for these activities",
            "No 18-discipline catalog or step-only wellness",
        ])
        add_bullet_slide(prs, "4VELO position", [
            "White-label multi-tenant",
            "4-layer anti-cheat + public trust policy",
            "Mobile resilience (outbox) + scale-tested backend 10k–300k",
            "Operator console + 1v1 leagues in API",
        ])
        add_bullet_slide(prs, "Trust and scale", [
            "Public Anti-Cheat / Scoring Policy (live URL)",
            "RSP-weekend runbook + 10k–300k simulation proof",
            "DPIA + per-tenant subprocessor list",
        ])
        add_bullet_slide(prs, "Three revenue paths", [
            "Cities / local government — seasonal campaigns",
            "Companies — department leagues / company vs company (HR)",
            "Clubs — club vs club, federation seasons",
        ])
        add_bullet_slide(prs, "Proof pilots (90 days)", [
            "Pilot A: 4 departments, 6 weeks, 1 tenant",
            "Pilot B: 2 clubs, 30 days, 1v1",
            "Then: multi-company league + case study",
        ])
        add_title_slide(prs, "Next step", "30-min demo · DPIA · pilot start within 2 weeks")
    return prs


def main() -> None:
    OUT_PL.parent.mkdir(parents=True, exist_ok=True)
    OUT_EN.parent.mkdir(parents=True, exist_ok=True)
    build_deck("pl").save(OUT_PL)
    build_deck("en").save(OUT_EN)
    print(f"Wrote {OUT_PL}")
    print(f"Wrote {OUT_EN}")


if __name__ == "__main__":
    main()
